"""
Government-ready compliance exports — PDF & Excel reports for DRC inspectors.
"""

from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone
from django.contrib.admin.views.decorators import staff_member_required

# PDF
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer,
)

# Excel
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from apps.experts.models import ExpertProfile
from apps.visas.models import (
    CarteDeTravail, EstablishmentVisa, AirportVisa, FlyingVisa, VoyageVisa,
)
from apps.compliance.models import PenaltyLedger, ExportLog


# ═══════════════════════════════════════════════════════════════
#  DATA COLLECTION
# ═══════════════════════════════════════════════════════════════

def _collect_data():
    """Gather compliance data for the current tenant's active experts."""
    today = timezone.now().date()
    experts = (
        ExpertProfile.objects
        .filter(is_active=True)
        .order_by("last_name", "first_name")
    )

    rows = []
    for e in experts:
        carte = (
            CarteDeTravail.objects
            .filter(expert=e, is_active=True)
            .order_by("-expiry_date")
            .first()
        )
        est_visa = (
            EstablishmentVisa.objects
            .filter(expert=e)
            .exclude(status__in=["CANCELLED", "EXPIRED"])
            .order_by("-expiry_date")
            .first()
        )
        penalties = PenaltyLedger.objects.filter(expert=e, resolved=False)
        penalty_total = sum(p.total_amount_usd for p in penalties)

        rows.append({
            "expert": e,
            "passport_expiry": e.passport_expiry_date,
            "passport_days": (e.passport_expiry_date - today).days if e.passport_expiry_date else None,
            "yellow_fever_expiry": e.yellow_fever_expiry_date,
            "carte_number": carte.permit_number if carte else "—",
            "carte_expiry": carte.expiry_date if carte else None,
            "carte_days": (carte.expiry_date - today).days if carte else None,
            "visa_number": est_visa.visa_number if est_visa else "—",
            "visa_expiry": est_visa.expiry_date if est_visa else None,
            "visa_days": (est_visa.expiry_date - today).days if est_visa else None,
            "penalty_total": float(penalty_total),
        })
    return rows, today


# ═══════════════════════════════════════════════════════════════
#  AUDIT LOGGING FOR EXPORTS
# ═══════════════════════════════════════════════════════════════

def _log_export(request, fmt, report_type, row_count, filters=None):
    """Record every export in the audit trail (SOC 2 requirement)."""
    try:
        ExportLog.objects.create(
            user_id=getattr(request.user, "id", None),
            user_repr=str(request.user) if request.user else "anonymous",
            format=fmt,
            report_type=report_type,
            row_count=row_count,
            filters=filters or {},
            ip_address=request.META.get("REMOTE_ADDR"),
            user_agent=request.META.get("HTTP_USER_AGENT", "")[:500],
        )
    except Exception:
        # Never let audit logging break a user-facing export
        import logging
        logging.getLogger(__name__).exception("Export logging failed")


def _status_for(row):
    """Derive overall compliance status for one expert row."""
    p = row["passport_days"]
    c = row["carte_days"]
    v = row["visa_days"]

    if (p is not None and p < 0) or row["penalty_total"] > 0:
        return "IRREGULAR"
    if (
        (c is not None and c < 30)
        or (v is not None and v < 30)
        or (p is not None and p < 30)
    ):
        return "CRITICAL"
    if (
        (p is not None and p < 90)
        or (c is not None and c < 90)
        or (v is not None and v < 90)
    ):
        return "ATTENTION"
    return "COMPLIANT"


# ═══════════════════════════════════════════════════════════════
#  PDF EXPORT
# ═══════════════════════════════════════════════════════════════

@staff_member_required
def export_pdf(request):
    rows, today = _collect_data()

    _log_export(request, ExportLog.Format.PDF, "Compliance Report", len(rows))

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=1.2 * cm,
        rightMargin=1.2 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.2 * cm,
        title="DRC Expat Compliance Report",
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TitleX", parent=styles["Title"], fontSize=18,
        textColor=colors.HexColor("#0B132B"), alignment=0, spaceAfter=4,
    )
    sub_style = ParagraphStyle(
        "SubX", parent=styles["Normal"], fontSize=10,
        textColor=colors.HexColor("#64748B"), alignment=0, spaceAfter=14,
    )
    cell_style = ParagraphStyle(
        "CellX", parent=styles["Normal"], fontSize=8, leading=10,
    )
    header_style = ParagraphStyle(
        "HeaderX", parent=styles["Normal"], fontSize=8, leading=10,
        textColor=colors.white, fontName="Helvetica-Bold",
    )

    elements = []
    elements.append(Paragraph("DRC EXPAT COMPLIANCE REPORT", title_style))
    elements.append(Paragraph(
        f"Rapport de Conformité DRC Expat &nbsp;|&nbsp; "
        f"Generated: {today.strftime('%B %d, %Y')} &nbsp;|&nbsp; "
        f"Active experts: {len(rows)}",
        sub_style,
    ))

    headers = [
        "#", "Expert", "Nationality", "Passport Expiry",
        "Carte de Travail", "Carte Expiry", "Visa Expiry",
        "Penalty (USD)", "Status",
    ]
    data = [[Paragraph(h, header_style) for h in headers]]

    for i, r in enumerate(rows, 1):
        p = r["passport_days"]
        c = r["carte_days"]
        v = r["visa_days"]

        passport_status = "EXPIRED" if (p is not None and p < 0) else (f"{p}d" if p is not None else "—")
        carte_status = "—" if c is None else ("EXPIRED" if c < 0 else f"{c}d")
        visa_status = "—" if v is None else ("EXPIRED" if v < 0 else f"{v}d")

        data.append([
            Paragraph(str(i), cell_style),
            Paragraph(f"{r['expert'].last_name}, {r['expert'].first_name}", cell_style),
            Paragraph(r["expert"].nationality, cell_style),
            Paragraph(f"{r['passport_expiry']} ({passport_status})", cell_style),
            Paragraph(r["carte_number"], cell_style),
            Paragraph(f"{r['carte_expiry']} ({carte_status})" if r["carte_expiry"] else "—", cell_style),
            Paragraph(f"{r['visa_expiry']} ({visa_status})" if r["visa_expiry"] else "—", cell_style),
            Paragraph(f"${r['penalty_total']:.2f}" if r["penalty_total"] else "—", cell_style),
            Paragraph(_status_for(r), cell_style),
        ])

    table = Table(
        data,
        colWidths=[1 * cm, 4 * cm, 2.5 * cm, 3.5 * cm, 3 * cm, 3 * cm, 3 * cm, 2.5 * cm, 2.5 * cm],
        repeatRows=1,
    )
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0B132B")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 8),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
        ("TOPPADDING", (0, 0), (-1, 0), 8),
        ("BACKGROUND", (0, 1), (-1, -1), colors.white),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E2E8F0")),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 1), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 1), (-1, -1), 6),
    ]))
    elements.append(table)

    elements.append(Spacer(1, 14))
    footer_style = ParagraphStyle(
        "FootX", parent=styles["Normal"], fontSize=8,
        textColor=colors.HexColor("#64748B"), alignment=1,
    )
    elements.append(Paragraph(
        "This report is generated by the DRC Expat Expert & Visa Compliance Platform. "
        "For official use by DGM, Ministry of Labour, and corporate compliance officers.",
        footer_style,
    ))

    doc.build(elements)
    buffer.seek(0)

    filename = f"DRC_Compliance_Report_{today.isoformat()}.pdf"
    response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


# ═══════════════════════════════════════════════════════════════
#  EXCEL EXPORT
# ═══════════════════════════════════════════════════════════════

@staff_member_required
def export_excel(request):
    rows, today = _collect_data()

    _log_export(request, ExportLog.Format.EXCEL, "Compliance Report", len(rows))

    wb = Workbook()
    ws = wb.active
    ws.title = "Compliance Report"

    header_fill = PatternFill(start_color="0B132B", end_color="0B132B", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    cell_font = Font(size=10)
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    ws.merge_cells("A1:I1")
    ws["A1"] = "DRC EXPAT COMPLIANCE REPORT"
    ws["A1"].font = Font(bold=True, size=16, color="0B132B")
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:I2")
    ws["A2"] = (
        f"Generated: {today.strftime('%B %d, %Y')} | "
        f"Active experts: {len(rows)} | Africa/Kinshasa"
    )
    ws["A2"].font = Font(size=10, color="64748B")

    headers = [
        "#", "Last Name", "First Name", "Nationality", "Passport Expiry",
        "Carte de Travail", "Carte Expiry", "Visa Expiry", "Penalty (USD)",
    ]
    ws.append([])
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=4, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[4].height = 22

    for i, r in enumerate(rows, 1):
        ws.append([
            i,
            r["expert"].last_name,
            r["expert"].first_name,
            r["expert"].nationality,
            str(r["passport_expiry"]) if r["passport_expiry"] else "—",
            r["carte_number"],
            str(r["carte_expiry"]) if r["carte_expiry"] else "—",
            str(r["visa_expiry"]) if r["visa_expiry"] else "—",
            f"${r['penalty_total']:.2f}" if r["penalty_total"] else "—",
        ])
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=ws.max_row, column=col)
            cell.font = cell_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="left", vertical="center")

        # Highlight passport column if expiring / expired
        pc = ws.cell(row=ws.max_row, column=5)
        if r["passport_days"] is not None:
            if r["passport_days"] < 0:
                pc.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
            elif r["passport_days"] < 90:
                pc.fill = PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid")

    for col, width in zip("ABCDEFGHI", [5, 18, 18, 14, 18, 20, 16, 16, 14]):
        ws.column_dimensions[col].width = width

    # Penalty summary sheet
    ws2 = wb.create_sheet("Penalty Ledger")
    ws2.append([
        "Expert", "Reason", "Start Date",
        "Days Accrued", "Daily Rate (USD)", "Total (USD)",
    ])
    for col in range(1, 7):
        c = ws2.cell(row=1, column=col)
        c.fill = header_fill
        c.font = header_font
    for p in (
        PenaltyLedger.objects
        .filter(resolved=False)
        .order_by("-accrual_start_date")
    ):
        ws2.append([
            p.expert.full_name,
            p.get_reason_display(),
            str(p.accrual_start_date),
            p.days_accrued,
            float(p.daily_rate_usd),
            float(p.total_amount_usd),
        ])
    for col, width in zip("ABCDEF", [22, 22, 14, 14, 18, 16]):
        ws2.column_dimensions[col].width = width

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"DRC_Compliance_Report_{today.isoformat()}.xlsx"
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response